"""Reproducible binary classification evaluation from matched CSV evidence."""
import csv
import io
from collections import Counter
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def read_rows(content: str, required: set[str], key: str) -> list[dict]:
    reader = csv.DictReader(io.StringIO(content.lstrip('\ufeff')))
    if not required.issubset(reader.fieldnames or []):
        raise ValueError(f"CSV requires columns: {', '.join(sorted(required))}")
    rows = list(reader)
    if not rows:
        raise ValueError('CSV must contain at least one data row.')
    seen = set()
    for row in rows:
        if None in row or any(row.get(field) is None for field in required):
            raise ValueError('Malformed CSV row.')
        for field in required:
            row[field] = row[field].strip()
        if not row[key] or row[key] in seen:
            raise ValueError(f'{key} must be nonempty and unique.')
        seen.add(row[key])
    return rows


def build_evaluation_report(predictions: str, reference: str, tests: str,
                            source: str, validated_by: str, scope: str) -> str:
    predicted = read_rows(predictions, {'unit_id', 'predicted_label'}, 'unit_id')
    actual = read_rows(reference, {'unit_id', 'reference_label'}, 'unit_id')
    cases = read_rows(tests, {'case_code', 'status'}, 'case_code')
    if source not in {'expert_validated', 'authorized_pnp', 'synthetic_demo'}:
        raise ValueError('Select a supported reference source.')
    if not validated_by.strip() or not scope.strip():
        raise ValueError('Reference validator/source and evaluation scope are required.')
    prediction_map = {r['unit_id']: r['predicted_label'] for r in predicted}
    reference_map = {r['unit_id']: r['reference_label'] for r in actual}
    if prediction_map.keys() != reference_map.keys():
        raise ValueError('Prediction and reference unit IDs must match exactly; include the complete evaluation population.')
    if any(v not in {'0', '1'} for v in [*prediction_map.values(), *reference_map.values()]):
        raise ValueError('Labels must be 0 (non-hotspot) or 1 (hotspot).')
    statuses = Counter(r['status'] for r in cases)
    if statuses.keys() - {'passed', 'failed', 'blocked', 'not_run'}:
        raise ValueError('Test status must be passed, failed, blocked, or not_run.')
    counts = Counter((prediction_map[k], reference_map[k]) for k in prediction_map)
    tp, fp, fn, tn = (counts[k] for k in [('1', '1'), ('1', '0'), ('0', '1'), ('0', '0')])
    # Explicit zero-division convention: a metric with no denominator is 0.
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * tp / (2 * tp + fp + fn) if 2 * tp + fp + fn else 0.0
    # Per the study methodology (SOP 4), the F1-Score alone is the Algorithmic
    # Performance Score, because it is the harmonic mean of Precision and
    # Recall. Precision and Recall are reported separately as supporting
    # measures and are NOT averaged into the composite score.
    algorithmic = f1
    functional = statuses['passed'] / len(cases)
    values = {
        'generated_at': datetime.now(timezone.utc).isoformat(),
        'reference_source': source,
        'reference_validator_or_source': validated_by.strip(),
        'evaluation_scope': scope.strip(),
        'validation_basis': 'User-supplied reference classification and provenance',
        'result_status': 'DEMONSTRATION ONLY - not validated PNP findings' if source == 'synthetic_demo' else 'Compared against user-supplied validated reference',
        'positive_class': '1 = hotspot; 0 = non-hotspot',
        'evaluation_units': len(predicted),
        'true_positives': tp, 'false_positives': fp,
        'false_negatives': fn, 'true_negatives': tn,
        'precision_percent': precision * 100, 'recall_percent': recall * 100,
        'f1_score_percent': f1 * 100,
        'zero_division_policy': 'Undefined ratios are reported as 0',
        'algorithmic_formula': 'Algorithmic Performance Score (%) = F1-Score x 100',
        'algorithmic_score_percent': algorithmic * 100,
        'algorithmic_weight_percent': 90,
        'functional_passed': statuses['passed'], 'functional_failed': statuses['failed'],
        'functional_blocked': statuses['blocked'], 'functional_not_run': statuses['not_run'],
        'functional_total': len(cases),
        'functional_formula': 'Passed / all supplied test cases (including blocked and not_run)',
        'functional_score_percent': functional * 100,
        'functional_weight_percent': 10,
        'algorithmic_contribution_percentage_points': algorithmic * 90,
        'functional_contribution_percentage_points': functional * 10,
        'overall_formula': '0.90 * algorithmic_score_percent + 0.10 * functional_score_percent',
        'overall_effectiveness_percent': algorithmic * 90 + functional * 10,
    }
    buffer = io.StringIO(newline='')
    writer = csv.writer(buffer)
    # Prevent spreadsheet formula execution in user-supplied provenance and IDs.
    def safe(value):
        if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
            return "'" + value
        return value
    writer.writerow(['section', 'item', 'value', 'reference_label'])
    for name, value in values.items():
        writer.writerow(['summary', name, safe(value), ''])
    for unit, label in prediction_map.items():
        writer.writerow(['classification', safe(unit), label, reference_map[unit]])
    for case in cases:
        writer.writerow(['functional_test', safe(case['case_code']), case['status'], ''])
    return buffer.getvalue()


def build_evaluation_workbook(predictions: str, reference: str, tests: str,
                              source: str, validated_by: str, scope: str) -> io.BytesIO:
    """Build a panel-ready workbook while retaining every supplied evidence row."""
    audit_csv = build_evaluation_report(
        predictions, reference, tests, source, validated_by, scope
    )
    rows = list(csv.DictReader(io.StringIO(audit_csv)))
    values = {r['item']: r['value'] for r in rows if r['section'] == 'summary'}
    classifications_data = [r for r in rows if r['section'] == 'classification']
    functional_data = [r for r in rows if r['section'] == 'functional_test']

    number = lambda key: float(values[key])
    count = lambda key: int(values[key])
    precision = number('precision_percent') / 100
    recall = number('recall_percent') / 100
    f1 = number('f1_score_percent') / 100
    functional_score = number('functional_score_percent') / 100
    overall = number('overall_effectiveness_percent') / 100
    tp, fp = count('true_positives'), count('false_positives')
    fn, tn = count('false_negatives'), count('true_negatives')
    status_counts = {
        name: count(f'functional_{name}')
        for name in ('passed', 'failed', 'blocked', 'not_run')
    }

    wb = Workbook()
    summary = wb.active
    summary.title = 'Summary'
    classifications = wb.create_sheet('Classifications')
    functional = wb.create_sheet('Functional tests')

    navy, slate, blue = '0F172A', '334155', '2563EB'
    pale_blue, white = 'EFF6FF', 'FFFFFF'
    amber, pale_amber = 'D97706', 'FFF7ED'
    green, pale_green = '15803D', 'F0FDF4'
    red, pale_red = 'B91C1C', 'FEF2F2'
    light_slate = 'E2E8F0'
    thin = Side(style='thin', color='CBD5E1')

    def safe(value):
        if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
            return "'" + value
        return value

    def add_title(sheet, text, end_column):
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
        cell = sheet.cell(2, 1, text)
        cell.font = Font(name='Arial', size=16, bold=True, color=navy)
        cell.alignment = Alignment(vertical='center')
        sheet.row_dimensions[2].height = 26

    def add_section(sheet, row, text, end_column):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_column)
        cell = sheet.cell(row, 1, text)
        cell.fill = PatternFill('solid', fgColor=navy)
        cell.font = Font(name='Arial', size=10, bold=True, color=white)
        cell.alignment = Alignment(vertical='center')
        sheet.row_dimensions[row].height = 20

    def add_headers(sheet, row, start_column, headings):
        for offset, heading in enumerate(headings):
            cell = sheet.cell(row, start_column + offset, heading)
            cell.fill = PatternFill('solid', fgColor=slate)
            cell.font = Font(name='Arial', size=10, bold=True, color=white)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(left=thin, right=thin)

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0

    summary.sheet_properties.tabColor = blue
    add_title(summary, 'PPALMA Hotspot Evaluation', 8)
    summary.merge_cells('A3:H3')
    summary['A3'] = values['result_status']
    summary['A3'].fill = PatternFill('solid', fgColor=pale_amber if source == 'synthetic_demo' else pale_blue)
    summary['A3'].font = Font(name='Arial', size=10, bold=True, color=amber if source == 'synthetic_demo' else blue)
    summary['A3'].alignment = Alignment(horizontal='center', vertical='center')
    summary.row_dimensions[3].height = 23

    add_section(summary, 5, 'Evaluation overview', 8)
    overview = [
        ('Generated (UTC)', values['generated_at'], 'Reference source', source.replace('_', ' ').title()),
        ('Validator / evidence', safe(values['reference_validator_or_source']), 'Evaluation units', count('evaluation_units')),
    ]
    for row_number, record in enumerate(overview, start=6):
        for column, value in zip((1, 2, 6, 7), record):
            summary.cell(row_number, column, value)
        summary.merge_cells(start_row=row_number, start_column=2, end_row=row_number, end_column=4)
        summary.merge_cells(start_row=row_number, start_column=7, end_row=row_number, end_column=8)
        summary.cell(row_number, 1).font = Font(name='Arial', size=10, bold=True, color=slate)
        summary.cell(row_number, 6).font = Font(name='Arial', size=10, bold=True, color=slate)
    summary['A8'] = 'Scope'
    summary['A8'].font = Font(name='Arial', size=10, bold=True, color=slate)
    summary.merge_cells('B8:H9')
    summary['B8'] = safe(values['evaluation_scope'])
    summary['B8'].alignment = Alignment(wrap_text=True, vertical='top')
    summary.row_dimensions[8].height = 32

    add_section(summary, 11, 'Weighted effectiveness score', 8)
    add_headers(summary, 12, 1, ['Component', 'Score', 'Weight', 'Contribution'])
    score_rows = [
        ('Algorithmic performance (F1)', f1, .90, f1 * 90),
        ('Functional pass rate', functional_score, .10, functional_score * 10),
        ('Overall effectiveness', overall, 1.0, overall * 100),
    ]
    for row_number, record in enumerate(score_rows, start=13):
        for column, value in enumerate(record, start=1):
            summary.cell(row_number, column, value)
        summary.cell(row_number, 2).number_format = '0.0%'
        summary.cell(row_number, 3).number_format = '0%'
        summary.cell(row_number, 4).number_format = '0.0 "pp"'
    for cell in summary[15][:4]:
        cell.fill = PatternFill('solid', fgColor=pale_blue)
        cell.font = Font(name='Arial', size=11, bold=True, color=navy)

    add_headers(summary, 12, 6, ['Classification metric', 'Result', 'Meaning'])
    for row_number, record in enumerate([
        ('Precision', precision, 'Correct among predicted hotspots'),
        ('Recall', recall, 'Detected among reference hotspots'),
        ('F1 score', f1, 'Algorithmic performance score'),
    ], start=13):
        for column, value in enumerate(record, start=6):
            summary.cell(row_number, column, value)
        summary.cell(row_number, 7).number_format = '0.0%'

    add_section(summary, 18, 'Confusion matrix and functional evidence', 8)
    add_headers(summary, 19, 1, ['', 'Predicted hotspot', 'Predicted non-hotspot', 'Total'])
    confusion_rows = [
        ('Reference hotspot', tp, fn, tp + fn),
        ('Reference non-hotspot', fp, tn, fp + tn),
        ('Total', tp + fp, fn + tn, count('evaluation_units')),
    ]
    for row_number, record in enumerate(confusion_rows, start=20):
        for column, value in enumerate(record, start=1):
            summary.cell(row_number, column, value)
    summary['B20'].fill = PatternFill('solid', fgColor=pale_green)
    summary['C20'].fill = PatternFill('solid', fgColor=pale_red)
    summary['B21'].fill = PatternFill('solid', fgColor=pale_red)
    summary['C21'].fill = PatternFill('solid', fgColor=pale_green)
    for cell in summary[22][:4]:
        cell.font = Font(name='Arial', size=10, bold=True, color=navy)

    add_headers(summary, 19, 6, ['Functional status', 'Cases', 'Share'])
    functional_total = count('functional_total')
    for row_number, name in enumerate(('passed', 'failed', 'blocked', 'not_run'), start=20):
        summary.cell(row_number, 6, name.replace('_', ' ').title())
        summary.cell(row_number, 7, status_counts[name])
        summary.cell(row_number, 8, status_counts[name] / functional_total)
        summary.cell(row_number, 8).number_format = '0.0%'
    summary['F24'], summary['G24'], summary['H24'] = 'Total', functional_total, 1.0
    summary['H24'].number_format = '0.0%'
    for cell in summary[24][5:8]:
        cell.font = Font(name='Arial', size=10, bold=True, color=navy)
    summary['F20'].font = Font(name='Arial', size=10, bold=True, color=green)

    add_section(summary, 27, 'Calculation method', 8)
    methods = [
        'Algorithmic Performance Score = F1 score. Precision and Recall are supporting metrics.',
        'Functional score = passed / all supplied cases. Failed, blocked, and not-run cases remain in the denominator.',
        'Overall effectiveness = (90% x algorithmic score) + (10% x functional score).',
        'Positive class: 1 = hotspot; 0 = non-hotspot. Undefined ratios are reported as 0.',
    ]
    for row_number, method in enumerate(methods, start=28):
        summary.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=8)
        summary.cell(row_number, 1, method)
        summary.cell(row_number, 1).alignment = Alignment(wrap_text=True, vertical='center')
    summary.merge_cells('A33:H33')
    summary['A33'] = ('Synthetic results demonstrate the evaluation workflow and are not validated PNP findings.'
                      if source == 'synthetic_demo' else values['validation_basis'])
    summary['A33'].fill = PatternFill('solid', fgColor=pale_amber if source == 'synthetic_demo' else pale_blue)
    summary['A33'].font = Font(name='Arial', size=9, italic=True, color=amber if source == 'synthetic_demo' else blue)
    summary['A33'].alignment = Alignment(horizontal='center')
    for column, width in {'A': 29, 'B': 19, 'C': 23, 'D': 16, 'E': 3, 'F': 25, 'G': 13, 'H': 34}.items():
        summary.column_dimensions[column].width = width
    summary.freeze_panes = 'A5'
    summary.print_area = 'A1:H33'
    summary.page_margins.left = .25
    summary.page_margins.right = .25

    add_title(classifications, 'Classification evidence', 4)
    if source == 'synthetic_demo':
        classifications.merge_cells('A3:D3')
        classifications['A3'] = 'Synthetic demonstration only'
        classifications['A3'].fill = PatternFill('solid', fgColor=pale_amber)
        classifications['A3'].font = Font(name='Arial', size=10, bold=True, color=amber)
        classifications['A3'].alignment = Alignment(horizontal='center')
    add_headers(classifications, 5, 1, ['Unit ID', 'System prediction', 'Reference classification', 'Result'])
    outcomes = {('1', '1'): 'True positive', ('1', '0'): 'False positive',
                ('0', '1'): 'False negative', ('0', '0'): 'True negative'}
    for row_number, record in enumerate(classifications_data, start=6):
        predicted, reference_label = record['value'], record['reference_label']
        classifications.append([
            safe(record['item']), int(predicted), int(reference_label),
            outcomes[(predicted, reference_label)],
        ])
        classifications.cell(row_number, 2).alignment = Alignment(horizontal='center')
        classifications.cell(row_number, 3).alignment = Alignment(horizontal='center')
        outcome = classifications.cell(row_number, 4)
        matched = predicted == reference_label
        outcome.fill = PatternFill('solid', fgColor=pale_green if matched else pale_red)
        outcome.font = Font(name='Arial', size=10, color=green if matched else red)
    classifications.freeze_panes = 'A6'
    classifications.auto_filter.ref = f'A5:D{5 + len(classifications_data)}'
    for column, width in {'A': 24, 'B': 20, 'C': 24, 'D': 20}.items():
        classifications.column_dimensions[column].width = width

    add_title(functional, 'Functional test evidence', 2)
    if source == 'synthetic_demo':
        functional.merge_cells('A3:B3')
        functional['A3'] = 'Synthetic demonstration only'
        functional['A3'].fill = PatternFill('solid', fgColor=pale_amber)
        functional['A3'].font = Font(name='Arial', size=10, bold=True, color=amber)
        functional['A3'].alignment = Alignment(horizontal='center')
    add_headers(functional, 5, 1, ['Case code', 'Status'])
    status_styles = {
        'passed': (pale_green, green), 'failed': (pale_red, red),
        'blocked': (pale_amber, amber), 'not_run': (light_slate, slate),
    }
    for row_number, record in enumerate(functional_data, start=6):
        status = record['value']
        functional.append([safe(record['item']), status.replace('_', ' ').title()])
        fill, color = status_styles[status]
        functional.cell(row_number, 2).fill = PatternFill('solid', fgColor=fill)
        functional.cell(row_number, 2).font = Font(name='Arial', size=10, color=color)
    functional.freeze_panes = 'A6'
    functional.auto_filter.ref = f'A5:B{5 + len(functional_data)}'
    functional.column_dimensions['A'].width = 28
    functional.column_dimensions['B'].width = 20

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font.name != 'Arial':
                    cell.font = Font(name='Arial', size=10, bold=cell.font.bold,
                                     italic=cell.font.italic, color=cell.font.color)
                if cell.value is not None and cell.alignment.vertical is None:
                    cell.alignment = Alignment(vertical='center')

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
