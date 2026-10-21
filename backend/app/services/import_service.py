"""Reusable CSV/Excel import module.

Deliberately generic: parse -> suggest column mapping -> validate -> commit.
It is exercised against synthetic CSVs today and is the same code path the
official PNP dataset will travel once its columns are known. Only the mapping
supplied at import time changes, not this module.

Validation outcomes follow the specification's three categories:

  valid    the row imports cleanly
  warning  the row imports, but something is worth the operator's attention
           (missing barangay, missing time, generalized coordinates)
  rejected the row cannot be imported without corrupting the analysis

Standardizations (for example "Pigkawayan" -> "Pigcawayan", or "Physical
Injuries" -> "Physical Injury") are applied but never silently: every one is
recorded and returned so the import summary can show exactly what changed.
"""
import csv
import io
import re
from collections import Counter
from datetime import date, datetime, time as time_type

import pandas as pd
from openpyxl import load_workbook
from geoalchemy2.shape import from_shape
from shapely.geometry import Point
from sqlalchemy.orm import Session

from app.core.taxonomy import (
    CRIMES_AGAINST_PERSONS,
    INDEX_CRIME_TYPES,
    normalize_crime_type,
    normalize_municipality,
)
from app.models.incident import Incident
from app.models.municipality import Municipality
from app.services.locality import barangay_centroid, canonical_barangay, nearest_barangay

KNOWN_CRIME_TYPES = set(INDEX_CRIME_TYPES)

TARGET_FIELDS = [
    "incident_id",
    "crime_type",
    "incident_date",
    "incident_time",
    "municipality",
    "barangay",
    "general_location",
    "latitude",
    "longitude",
    "coordinate_source",
    "coordinate_precision",
    "reporting_station",
    "case_status_generalized",
    "offense_code",
    "notes_non_sensitive",
]

# Latitude/longitude are deliberately absent: a row with a resolvable barangay
# is placed at that barangay's official centroid rather than dropped, since the
# PNP data-sharing agreement allows releasing barangay-level location when exact
# coordinates cannot be given. Rows are only rejected when neither is usable.
REQUIRED_FIELDS = {
    "incident_id",
    "crime_type",
    "incident_date",
    "municipality",
}

# Inbound header variants the mapping suggester recognizes automatically.
COLUMN_SYNONYMS: dict[str, list[str]] = {
    "incident_id": [
        "incident_id", "id", "case_id", "case_number", "case_no", "blotter_no",
        "blotterno",
        "blotter_number", "blotter_entry_no", "entry_no", "reference_no",
        "reference_number", "ref_no", "control_no",
    ],
    "crime_type": [
        "crime_type", "crime", "offense", "offense_type", "incident_type",
        "crime_category", "type_of_crime", "nature_of_offense", "nature_of_incident",
        "offense_committed", "crime_committed",
    ],
    "incident_date": [
        "incident_date", "date", "date_committed", "date_of_incident",
        "date_reported", "date_time_committed", "datetime", "date_of_commission",
    ],
    "incident_time": [
        "incident_time", "time", "time_committed", "time_of_incident",
        "time_reported", "approximate_time",
    ],
    "municipality": [
        "municipality", "municipal", "city_municipality", "city_mun", "citymun", "lgu", "town",
        "municipality_city", "mun", "city",
    ],
    "barangay": ["barangay", "brgy", "bgy", "barangay_name", "brgy_name", "village"],
    "general_location": [
        "general_location", "location", "location_description", "place",
        "address_general", "place_of_incident", "place_of_occurrence",
        "location_of_incident", "sitio_purok", "landmark", "street",
    ],
    "latitude": ["latitude", "lat", "y", "ycoord", "y_coordinate", "gps_latitude"],
    "longitude": ["longitude", "lon", "lng", "long", "x", "xcoord", "x_coordinate", "gps_longitude"],
    "coordinate_source": ["coordinate_source", "coord_source", "geocode_source"],
    "coordinate_precision": ["coordinate_precision", "coord_precision", "precision"],
    "reporting_station": [
        "reporting_station", "station", "police_station", "unit", "ps",
        "reporting_unit", "responding_unit",
    ],
    "case_status_generalized": [
        "case_status_generalized", "case_status", "status", "disposition", "case_disposition",
    ],
    "offense_code": ["offense_code", "code", "ra_number", "ra_no", "article_no"],
    "notes_non_sensitive": ["notes_non_sensitive", "notes", "remarks", "narrative_summary"],
}

# Reasonable envelope around the PPALMA study area. Coordinates outside this
# are almost certainly transposed or mis-entered; point-in-polygon still runs
# separately and is the authoritative check.
STUDY_AREA_BOUNDS = {"min_lat": 6.4, "max_lat": 7.9, "min_lon": 124.0, "max_lon": 125.2}

# Reserved column the reader adds to every worksheet row, naming the tab it
# came from. Prefixed so it cannot collide with a real exported heading.
SHEET_COLUMN = "__source_sheet__"

# Consecutive blank worksheet rows that mean "the records have ended".
BLANK_ROW_RUN_ENDS_SHEET = 500

MIN_YEAR = 2000
MAX_YEAR = date.today().year + 1


def _normalize_header(value: object) -> str:
    """Collapse a header cell to its synonym-lookup key.

    camelCase is split first, so the `dateCommitted` / `timeReported` style of
    a system export resolves to the same key as a `Date Committed` heading.
    """
    text = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", "_", str(value).strip())
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


_ALL_SYNONYMS = {s for synonyms in COLUMN_SYNONYMS.values() for s in synonyms}


def _find_header_row(raw: pd.DataFrame) -> int:
    """Index of the row that actually holds the column headers.

    Government worksheet exports routinely open with a title block ("Republic
    of the Philippines", a unit name, a blank line) before the real header, in
    which case row 0 is not it. The header is the row with the most cells that
    name a known field, falling back to the most populated row.
    """
    best_row, best_score = 0, -1
    for idx in range(min(len(raw), 15)):
        values = [
            str(v).strip()
            for v in raw.iloc[idx]
            if v is not None and str(v).strip() and str(v).strip().lower() != "nan"
        ]
        if len(values) < 2:
            continue
        recognized = sum(1 for v in values if _normalize_header(v) in _ALL_SYNONYMS)
        score = recognized * 100 + len(values)
        if score > best_score:
            best_row, best_score = idx, score
    return best_row


def _jsonable(value):
    """Render date/time cells as ISO text.

    Excel stores dates as real datetime cells, which would otherwise reach the
    validator as epoch milliseconds once the batch is serialized to JSON and be
    read as 1970. Converting here keeps every code path — CSV or worksheet —
    handing the validator the same unambiguous ISO string.
    """
    if isinstance(value, pd.Timestamp):
        return value.date().isoformat() if value == value.normalize() else value.isoformat()
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time_type() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time_type):
        return value.strftime("%H:%M:%S")
    return value


def _frame_from_raw(raw: pd.DataFrame) -> pd.DataFrame:
    """Turn a header-less worksheet read into a framed table."""
    header_index = _find_header_row(raw)
    header = [str(c).strip() for c in raw.iloc[header_index]]

    seen: dict[str, int] = {}
    columns = []
    for name in header:
        if not name or name.lower() == "nan":
            name = ""
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        columns.append(name)

    frame = raw.iloc[header_index + 1 :].copy()
    frame.columns = columns
    frame = frame.loc[:, [c for c in columns if c]]
    return frame.dropna(how="all").reset_index(drop=True)


def _decode(content: bytes) -> str:
    """Decode text exported by Windows tooling, which is often not UTF-8."""
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")  # single-byte: decodes anything


def _read_delimited(content: bytes) -> pd.DataFrame:
    """Read delimited text of any separator, encoding, or row width.

    Parsed row-by-row rather than through pandas' strict reader because a real
    export is ragged: a title block above the header has fewer fields than the
    table, which is a hard error for the strict parser but simply a short row
    here. The separator is chosen as the one producing the most rows of
    consistent width, so comma, semicolon, tab and pipe all work.
    """
    text = _decode(content)
    best: tuple[int, int, list[list[str]]] = (0, 0, [])
    for delimiter in (",", ";", "\t", "|"):
        try:
            rows = [r for r in csv.reader(io.StringIO(text), delimiter=delimiter) if any(c.strip() for c in r)]
        except csv.Error:
            continue
        if not rows:
            continue
        widths = Counter(len(r) for r in rows)
        width, count = widths.most_common(1)[0]
        if width > 1 and (count, width) > (best[1], best[0]):
            best = (width, count, rows)

    width, _, rows = best
    if not rows:
        raise ValueError("No delimited rows found (tried comma, semicolon, tab, pipe).")
    return pd.DataFrame([r + [None] * (width - len(r)) for r in rows if len(r) <= width])


def _read_xlsx_streaming(content: bytes) -> dict[str, pd.DataFrame]:
    """Read an .xlsx sheet by sheet, skipping rows that hold no values.

    Worksheets carry formatting far below their data, so a sheet of 1,500
    records can declare a million rows. Reading through pandas would
    materialize every one of those empty rows; streaming with openpyxl and
    dropping blank rows as they arrive keeps the cost proportional to the
    actual data.
    """
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheets: dict[str, pd.DataFrame] = {}
        for name in workbook.sheetnames:
            rows: list[list] = []
            blank_run = 0
            for row in workbook[name].iter_rows(values_only=True):
                if any(cell is not None and str(cell).strip() for cell in row):
                    rows.append(list(row))
                    blank_run = 0
                    continue
                blank_run += 1
                # Formatting applied to empty cells can extend a sheet to a
                # million rows. A real table does not contain a gap this wide,
                # so a long blank run means the records have ended; stopping
                # here avoids parsing the empty remainder of the sheet.
                if blank_run >= BLANK_ROW_RUN_ENDS_SHEET:
                    break
            if rows:
                width = max(len(r) for r in rows)
                sheets[name] = pd.DataFrame(
                    [r + [None] * (width - len(r)) for r in rows], dtype=object
                )
        return sheets
    finally:
        workbook.close()


def _read_sheets(content: bytes) -> dict[str, pd.DataFrame]:
    """Read a workbook into {sheet name: raw header-less frame}.

    The extension is not proof of the format. Record systems routinely emit an
    HTML table or a delimited text file under a .xls name, and a real .xlsx is
    just as often saved as .xls, so every reader that could apply is tried
    before giving up.
    """
    attempts: list[str] = []
    try:
        sheets = _read_xlsx_streaming(content)
        if sheets:
            return sheets
        attempts.append("openpyxl: no non-empty sheets")
    except Exception as exc:
        attempts.append(f"openpyxl: {exc}")

    for engine in ("xlrd",):
        try:
            return pd.read_excel(
                io.BytesIO(content), sheet_name=None, header=None, engine=engine
            )
        except Exception as exc:
            attempts.append(f"{engine}: {exc}")

    try:
        tables = pd.read_html(io.BytesIO(content), header=None)
        if tables:
            return {f"Table {i + 1}": table for i, table in enumerate(tables)}
        attempts.append("html: no tables found")
    except Exception as exc:
        attempts.append(f"html: {exc}")

    try:
        return {"Sheet1": _read_delimited(content)}
    except Exception as exc:
        attempts.append(f"delimited text: {exc}")

    raise ValueError(
        "Unreadable as a workbook, an HTML table, or delimited text. "
        "If the file opens in Excel, re-save it as .xlsx or CSV and try again. "
        f"Readers tried — {' | '.join(attempts)[:600]}"
    )


def _require_rows(frame: pd.DataFrame) -> pd.DataFrame:
    """Guard against a reader that 'succeeded' on bytes that are not a table."""
    if frame.empty or not len(frame.columns):
        raise ValueError(
            "No data rows could be read. Confirm the file is a spreadsheet or "
            "delimited export and that it has a header row above the records."
        )
    return frame


def parse_file(filename: str, content: bytes) -> tuple[pd.DataFrame, list[str]]:
    """Return (rows, notes). Notes describe what the reader had to work out."""
    if not filename.lower().endswith((".xlsx", ".xls")):
        # A CSV export can carry a title block above the header just as a
        # worksheet can, so it goes through the same framing.
        return _require_rows(_frame_from_raw(_read_delimited(content))).map(_jsonable), []

    sheets = _read_sheets(content)
    notes: list[str] = []
    framed = {
        name: _frame_from_raw(raw)
        for name, raw in sheets.items()
        if not raw.dropna(how="all").empty
    }
    if not framed:
        raise ValueError("The workbook contains no data rows.")

    def recognized(frame: pd.DataFrame) -> set[str]:
        return {_normalize_header(c) for c in frame.columns} & _ALL_SYNONYMS

    primary_name, primary = next(iter(framed.items()))
    primary_fields = recognized(primary)
    combined = {primary_name: primary}
    skipped = []
    for name, frame in list(framed.items())[1:]:
        # Per-year tabs belong to one dataset even when their headings drift
        # between years — a date column renamed, an extra column added in a
        # single year. They are matched on the fields they have in common and
        # aligned, so no year is dropped for a cosmetic difference. A cover
        # sheet or legend shares nothing and is left out instead.
        if len(recognized(frame) & primary_fields) >= 2:
            combined[name] = frame
        else:
            skipped.append(name)

    for name, frame in combined.items():
        # Carries the tab's own name onto every row, so a sheet that identifies
        # its period only in the tab title ("CY 2018") can still be dated.
        frame[SHEET_COLUMN] = name

    if len(sheets) > 1:
        notes.append(f"Worksheets combined: {', '.join(combined)}.")
        if skipped:
            notes.append(
                f"Worksheets skipped (no columns in common): {', '.join(skipped)}."
            )
        widths = {tuple(f.columns) for f in combined.values()}
        if len(widths) > 1:
            notes.append(
                "Worksheet columns differed between tabs; rows were aligned by "
                "column name and missing columns left blank."
            )

    frames = list(combined.values())
    result = pd.concat(frames, ignore_index=True) if len(frames) > 1 else frames[0]
    return _require_rows(result).map(_jsonable), notes


def suggest_mapping(columns: list[str]) -> dict[str, str]:
    """Best-effort automatic column mapping; the operator can always override."""
    # Any run of punctuation or spacing collapses to one underscore, so headers
    # like "City/Municipality" or "Blotter Entry No." match their synonyms.
    normalized = {_normalize_header(c): c for c in columns}
    mapping: dict[str, str] = {}
    for field, synonyms in COLUMN_SYNONYMS.items():
        for candidate in synonyms:
            if candidate in normalized:
                mapping[field] = normalized[candidate]
                break
    return mapping


def _clean(raw):
    """Normalize pandas' empty representations to None."""
    if raw is None:
        return None
    if isinstance(raw, float) and raw != raw:  # NaN
        return None
    if isinstance(raw, str):
        stripped = raw.strip()
        return stripped or None
    return raw


_NUMERIC_DATE = re.compile(r"^\s*(\d{1,2})[/-](\d{1,2})[/-]\d{2,4}")


def _detect_dayfirst(rows: list[dict], column: str | None) -> tuple[bool, bool]:
    """Decide day-first vs month-first from the column's own values.

    Returns (dayfirst, proven). A leading component above 12 can only be a day;
    a second component above 12 can only mean month-first. Blotter exports are
    frequently D/M/Y while pandas assumes M/D/Y, so every ambiguous value (both
    parts <= 12) would otherwise be silently transposed into the wrong month.
    When no value in the column settles the order, `proven` is False and the
    caller flags the ambiguous rows instead of quietly picking one reading.
    """
    if not column:
        return False, True
    day_first = month_first = 0
    for row in rows:
        match = _NUMERIC_DATE.match(str(row.get(column) or ""))
        if not match:
            continue
        first, second = int(match.group(1)), int(match.group(2))
        if first > 12:
            day_first += 1
        elif second > 12:
            month_first += 1
    return day_first > month_first, bool(day_first or month_first)


def _is_ambiguous_date(raw: object) -> bool:
    """True for a numeric date whose day/month order the value cannot settle."""
    match = _NUMERIC_DATE.match(str(raw))
    return bool(match) and int(match.group(1)) <= 12 and int(match.group(2)) <= 12


_ISO_DATE = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[T ].*)?$")
_SLASH_DATE = re.compile(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$")
_CLOCK = re.compile(r"^(\d{1,2}):(\d{2})(?::\d{2})?(?:\.\d+)?$")


def _parse_date(raw: object, dayfirst: bool) -> date:
    """Read the formats real exports contain; fall back to pandas for the rest.

    pandas' scalar `to_datetime` costs roughly 300us, which on a six-figure
    import is minutes of pure parsing. These paths are three orders of
    magnitude cheaper and cover every value a spreadsheet or blotter produces;
    anything unusual still goes to pandas, so nothing that parsed before stops
    parsing now.
    """
    text = str(raw).strip()
    iso = _ISO_DATE.match(text)
    if iso:
        return date(int(iso[1]), int(iso[2]), int(iso[3]))

    slash = _SLASH_DATE.match(text)
    if slash:
        first, second, year = int(slash[1]), int(slash[2]), int(slash[3])
        if year < 100:
            year += 2000 if year < 70 else 1900
        day, month = (first, second) if dayfirst else (second, first)
        try:
            return date(year, month, day)
        except ValueError:
            pass  # the assumed order is impossible — let pandas resolve it

    return pd.to_datetime(text, dayfirst=dayfirst).date()


_SHEET_YEAR = re.compile(r"(19|20)\d{2}")


def _sheet_year(sheet_name: object) -> int | None:
    """Year named by a worksheet tab, e.g. 'CY 2018' -> 2018."""
    match = _SHEET_YEAR.search(str(sheet_name or ""))
    if not match:
        return None
    year = int(match.group(0))
    return year if MIN_YEAR <= year <= MAX_YEAR else None


def _parse_clock(raw: object) -> str:
    """Return HH:MM for a time cell."""
    text = str(raw).strip()
    clock = _CLOCK.match(text)
    if clock:
        hour, minute = int(clock[1]), int(clock[2])
        if hour < 24 and minute < 60:
            return f"{hour:02d}:{minute:02d}"
    return pd.to_datetime(text).time().strftime("%H:%M")


def validate_rows(
    rows: list[dict],
    mapping: dict[str, str],
    known_municipalities: set[str],
    id_prefix: str = "AUTO",
) -> tuple[list[dict], list[dict], list[dict]]:
    """Return (valid_rows, issues, standardizations).

    `issues` entries carry a `severity` of "warning" or "rejected"; a row is
    dropped only for a "rejected" issue.

    The importer does the organizing an operator would otherwise do by hand:
    barangay spellings are resolved to their official PSGC names, a missing
    barangay is recovered from the coordinate, a missing coordinate is recovered
    from the barangay centroid, date columns are read in the order the file
    itself proves, and a missing reference number is assigned one. Every such
    fix is reported as a standardization or warning — none are silent.
    """
    valid_rows: list[dict] = []
    issues: list[dict] = []
    standardizations: list[dict] = []
    seen_incident_ids: dict[str, int] = {}
    dayfirst, date_order_proven = _detect_dayfirst(rows, mapping.get("incident_date"))

    def reject(row_num: int, field: str, message: str) -> None:
        issues.append({"row": row_num, "field": field, "severity": "rejected", "message": message})

    def warn(row_num: int, field: str, message: str) -> None:
        issues.append({"row": row_num, "field": field, "severity": "warning", "message": message})

    def standardize(row_num: int, field: str, before, after) -> None:
        standardizations.append(
            {"row": row_num, "field": field, "from": str(before), "to": str(after)}
        )

    for idx, row in enumerate(rows):
        row_num = idx + 1
        parsed: dict = {}
        row_ok = True

        for field in TARGET_FIELDS:
            source_col = mapping.get(field)
            raw = _clean(row.get(source_col)) if source_col else None

            if raw is None and field == "incident_date":
                # A tab that names its period only in the title ("CY 2018") is
                # still usable: the study's periods are years, so the year is
                # recovered and the row is marked year-only precision. The
                # placeholder day is never presented as the day of the offense.
                year = _sheet_year(row.get(SHEET_COLUMN))
                if year is None:
                    reject(
                        row_num,
                        field,
                        "No incident date, and the worksheet name carries no year to fall back on",
                    )
                    row_ok = False
                    parsed[field] = None
                    continue
                parsed[field] = date(year, 1, 1).isoformat()
                parsed["date_precision"] = "year_only"
                standardize(row_num, field, f"(no date; {row.get(SHEET_COLUMN)})", f"{year} — year only")
                continue

            if raw is None:
                if field == "incident_id":
                    # A blotter export missing its reference number still holds a
                    # usable incident; assign a deterministic id (stable across
                    # re-validation of the same batch) instead of dropping it.
                    generated = f"{id_prefix}-{row_num:06d}"
                    standardize(row_num, field, "(blank)", generated)
                    parsed[field] = generated
                    continue
                if field in REQUIRED_FIELDS:
                    reject(row_num, field, "Required field is missing")
                    row_ok = False
                elif field == "incident_time":
                    warn(row_num, field, "Incident time is missing — date-only analysis for this row")
                parsed[field] = None
                continue

            if field == "incident_date":
                try:
                    parsed_date = _parse_date(raw, dayfirst)
                except Exception:
                    reject(row_num, field, f"Unparseable date: {raw}")
                    row_ok = False
                    continue
                if not (MIN_YEAR <= parsed_date.year <= MAX_YEAR):
                    reject(
                        row_num,
                        field,
                        f"Date {parsed_date.isoformat()} falls outside the plausible range "
                        f"{MIN_YEAR}-{MAX_YEAR}",
                    )
                    row_ok = False
                    continue
                if not date_order_proven and _is_ambiguous_date(raw):
                    reading = "day/month" if dayfirst else "month/day"
                    warn(
                        row_num,
                        field,
                        f"'{raw}' is ambiguous and nothing in this column settles the "
                        f"order — read as {reading} ({parsed_date.isoformat()}). Verify "
                        "against the source before relying on this row's date.",
                    )
                parsed[field] = parsed_date.isoformat()
                continue

            if field == "incident_time":
                try:
                    parsed[field] = _parse_clock(raw)
                except Exception:
                    warn(row_num, field, f"Unreadable time '{raw}' — imported without a time")
                    parsed[field] = None
                continue

            if field in ("latitude", "longitude"):
                try:
                    parsed[field] = float(raw)
                except (TypeError, ValueError):
                    # Left unset so barangay-centroid placement can recover the
                    # row; only a row with no usable location at all is rejected.
                    warn(row_num, field, f"Not a number: {raw}")
                    parsed[field] = None
                continue

            if field == "municipality":
                canonical, changed = normalize_municipality(raw)
                if canonical is None:
                    reject(
                        row_num,
                        field,
                        f"'{raw}' is not one of the six PPALMA municipalities "
                        f"({', '.join(sorted(known_municipalities))})",
                    )
                    row_ok = False
                    continue
                if changed:
                    standardize(row_num, field, raw, canonical)
                parsed[field] = canonical
                continue

            if field == "crime_type":
                canonical, changed = normalize_crime_type(raw)
                if canonical is None:
                    reject(
                        row_num,
                        field,
                        f"'{raw}' is not one of the eight index crimes analyzed by this study "
                        f"({', '.join(INDEX_CRIME_TYPES)})",
                    )
                    row_ok = False
                    continue
                if changed:
                    standardize(row_num, field, raw, canonical)
                parsed[field] = canonical
                continue

            parsed[field] = raw

        if not row_ok:
            continue

        municipality = parsed["municipality"]

        # --- Barangay: resolve the inbound spelling to its official PSGC name ---
        raw_barangay = parsed.get("barangay")
        if raw_barangay:
            official = canonical_barangay(municipality, raw_barangay)
            if official is None:
                warn(
                    row_num,
                    "barangay",
                    f"'{raw_barangay}' does not match any listed {municipality} barangay",
                )
            elif official != str(raw_barangay).strip():
                standardize(row_num, "barangay", raw_barangay, official)
            parsed["barangay"] = official

        lat, lon = parsed.get("latitude"), parsed.get("longitude")
        if lat is not None and not (-90 <= lat <= 90):
            warn(row_num, "latitude", f"Out of valid range: {lat}")
            lat = None
        if lon is not None and not (-180 <= lon <= 180):
            warn(row_num, "longitude", f"Out of valid range: {lon}")
            lon = None

        # --- Coordinate: fall back to the barangay's official centroid ---
        if lat is None or lon is None:
            centroid = (
                barangay_centroid(municipality, parsed["barangay"])
                if parsed.get("barangay")
                else None
            )
            if centroid is None:
                reject(
                    row_num,
                    "latitude/longitude",
                    "No usable location: the row has no valid coordinate and no "
                    "barangay that can be placed on the map",
                )
                continue
            lat, lon = centroid
            parsed["coordinate_source"] = "barangay_centroid"
            parsed["coordinate_precision"] = "barangay_level"
            warn(
                row_num,
                "latitude/longitude",
                f"Placed at the {parsed['barangay']} barangay centroid — "
                "cluster placement for this row is barangay-level, not exact",
            )
        parsed["latitude"], parsed["longitude"] = lat, lon

        # --- Barangay recovered from the coordinate when the cell was unusable ---
        if not parsed.get("barangay"):
            nearest = nearest_barangay(municipality, lat, lon)
            if nearest:
                standardize(row_num, "barangay", raw_barangay or "(blank)", nearest)
                parsed["barangay"] = nearest
            else:
                warn(row_num, "barangay", "Barangay is missing — imported without barangay detail")

        # Provenance defaults, so every stored row states how it was located
        # and how precisely it is dated.
        if not parsed.get("coordinate_source"):
            parsed["coordinate_source"] = "reported"
        if not parsed.get("coordinate_precision"):
            parsed["coordinate_precision"] = "exact"
        if not parsed.get("date_precision"):
            parsed["date_precision"] = "exact"

        outside = not (
            STUDY_AREA_BOUNDS["min_lat"] <= lat <= STUDY_AREA_BOUNDS["max_lat"]
            and STUDY_AREA_BOUNDS["min_lon"] <= lon <= STUDY_AREA_BOUNDS["max_lon"]
        )
        if outside:
            # A transposed or mis-keyed coordinate still has a usable barangay
            # most of the time; fall back rather than lose the incident.
            centroid = (
                barangay_centroid(municipality, parsed["barangay"])
                if parsed.get("barangay")
                else None
            )
            if centroid is None:
                reject(
                    row_num,
                    "latitude/longitude",
                    f"Coordinate ({lat}, {lon}) falls well outside the PPALMA study area — "
                    f"check for transposed latitude/longitude values",
                )
                continue
            parsed["latitude"], parsed["longitude"] = centroid
            parsed["coordinate_source"] = "barangay_centroid"
            parsed["coordinate_precision"] = "barangay_level"
            warn(
                row_num,
                "latitude/longitude",
                f"Coordinate ({lat}, {lon}) falls outside the PPALMA study area — "
                f"replaced with the {parsed['barangay']} barangay centroid",
            )
        elif parsed["coordinate_source"] != "barangay_centroid":
            precision = (parsed.get("coordinate_precision") or "").lower()
            if precision in ("barangay_level", "generalized", "approximate"):
                warn(
                    row_num,
                    "coordinate_precision",
                    f"Coordinate is '{precision}' — cluster placement for this row is approximate",
                )

        incident_id = str(parsed["incident_id"])
        if incident_id in seen_incident_ids:
            reject(
                row_num,
                "incident_id",
                f"Duplicate incident_id within this file (first seen at row "
                f"{seen_incident_ids[incident_id]})",
            )
            continue
        seen_incident_ids[incident_id] = row_num

        valid_rows.append(parsed)

    return valid_rows, issues, standardizations


def commit_rows(
    db: Session,
    valid_rows: list[dict],
    source: str = "official",
    dataset_id: int | None = None,
) -> tuple[int, int]:
    """Insert validated rows. Returns (inserted, skipped_existing)."""
    municipalities = {m.name: m.id for m in db.query(Municipality).all()}
    inserted = 0
    skipped = 0

    # Which ids are already stored, resolved in a handful of queries rather than
    # one per row — a full-year PNP export would otherwise issue a round trip
    # per record and take minutes instead of seconds.
    incoming = [row["incident_id"] for row in valid_rows]
    existing: set[str] = set()
    for start in range(0, len(incoming), 5000):
        chunk = incoming[start : start + 5000]
        existing.update(
            found
            for (found,) in db.query(Incident.incident_id)
            .filter(Incident.incident_id.in_(chunk))
            .all()
        )

    for row in valid_rows:
        if row["incident_id"] in existing:
            skipped += 1
            continue
        existing.add(row["incident_id"])

        incident_time = None
        if row.get("incident_time"):
            incident_time = datetime.strptime(row["incident_time"], "%H:%M").time()

        crime_type = row["crime_type"]
        point = Point(row["longitude"], row["latitude"])
        db.add(
            Incident(
                incident_id=row["incident_id"],
                crime_type=crime_type,
                incident_date=datetime.strptime(row["incident_date"], "%Y-%m-%d").date(),
                incident_time=incident_time,
                municipality_id=municipalities[row["municipality"]],
                municipality=row["municipality"],
                barangay=row.get("barangay") or "Unspecified",
                general_location=row.get("general_location"),
                latitude=row["latitude"],
                longitude=row["longitude"],
                geom=from_shape(point, srid=4326),
                crime_classification=(
                    "Crime Against Person"
                    if crime_type in CRIMES_AGAINST_PERSONS
                    else "Crime Against Property"
                ),
                index_nonindex_flag="index",
                coordinate_source=row.get("coordinate_source"),
                coordinate_precision=row.get("coordinate_precision"),
                date_precision=row.get("date_precision") or "exact",
                reporting_station=row.get("reporting_station"),
                case_status_generalized=row.get("case_status_generalized"),
                offense_code=row.get("offense_code"),
                notes_non_sensitive=row.get("notes_non_sensitive"),
                dataset_id=dataset_id,
                source=source,
            )
        )
        inserted += 1

    db.commit()
    return inserted, skipped
