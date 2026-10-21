import csv
import io
import unittest

from openpyxl import load_workbook

from app.services.evaluation_report import build_evaluation_report, build_evaluation_workbook


class EvaluationReportTests(unittest.TestCase):
    def report(self, predictions=None, reference=None, tests=None, source='expert_validated'):
        return build_evaluation_report(
            predictions if predictions is not None else 'unit_id,predicted_label\na,1\nb,1\nc,0\nd,0\n',
            reference if reference is not None else 'unit_id,reference_label\nd,0\nc,1\nb,0\na,1\n',
            tests if tests is not None else 'case_code,status\nT1,passed\nT2,failed\nT3,blocked\nT4,not_run\n',
            source, 'Expert review document 1', 'Run 1; complete grid; 2024',
        )

    def summary(self, report):
        return {row['item']: row['value'] for row in csv.DictReader(io.StringIO(report)) if row['section'] == 'summary'}

    def test_matched_ids_and_weighting(self):
        values = self.summary(self.report())
        for metric in ['precision_percent', 'recall_percent', 'f1_score_percent']:
            self.assertEqual(float(values[metric]), 50)
        self.assertEqual(float(values['functional_score_percent']), 25)
        self.assertEqual(float(values['overall_effectiveness_percent']), 47.5)
        for count in ['true_positives', 'false_positives', 'false_negatives', 'true_negatives']:
            self.assertEqual(values[count], '1')

    def test_algorithmic_score_is_f1_alone(self):
        """SOP 4: the Algorithmic Performance Score is F1 x 100 — not the mean
        of Precision, Recall and F1. This fixture is deliberately asymmetric
        (Precision 33.3%, Recall 100%, F1 50%) so the two formulas disagree:
        the mean would be 61.1%, F1 alone is 50%."""
        values = self.summary(
            self.report(
                predictions='unit_id,predicted_label\na,1\nb,1\nc,1\nd,0\n',
                reference='unit_id,reference_label\na,1\nb,0\nc,0\nd,0\n',
                tests='case_code,status\nT1,passed\n',
            )
        )
        self.assertAlmostEqual(float(values['precision_percent']), 100 / 3, places=6)
        self.assertEqual(float(values['recall_percent']), 100)
        self.assertEqual(float(values['f1_score_percent']), 50)
        self.assertEqual(float(values['algorithmic_score_percent']), 50)
        # 0.90 * 50 + 0.10 * 100
        self.assertEqual(float(values['overall_effectiveness_percent']), 55)

    def test_perfect_result(self):
        values = self.summary(self.report(reference='unit_id,reference_label\na,1\nb,1\nc,0\nd,0', tests='case_code,status\nT1,passed'))
        self.assertEqual(float(values['overall_effectiveness_percent']), 100)

    def test_zero_denominators(self):
        values = self.summary(self.report(predictions='unit_id,predicted_label\na,0', reference='unit_id,reference_label\na,0'))
        self.assertEqual(float(values['f1_score_percent']), 0)

    def test_rejects_invalid_evidence(self):
        for kwargs in [
            {'reference': 'unit_id,reference_label\na,1'},
            {'predictions': 'unit_id,predicted_label\na,1\na,0'},
            {'predictions': 'unit_id,predicted_label\na,2\nb,1\nc,0\nd,0'},
            {'tests': 'case_code,status\nT1,skipped'},
            {'tests': 'case_code,status\n'},
            {'reference': ''},
            {'source': 'unknown'},
        ]:
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                self.report(**kwargs)

    def test_demo_label_and_evidence_export(self):
        report = self.report(source='synthetic_demo')
        self.assertIn('DEMONSTRATION ONLY', report)
        rows = list(csv.DictReader(io.StringIO(report)))
        self.assertEqual(sum(r['section'] == 'classification' for r in rows), 4)
        self.assertEqual(sum(r['section'] == 'functional_test' for r in rows), 4)

    def test_panel_workbook_has_summary_and_evidence_sheets(self):
        workbook = load_workbook(build_evaluation_workbook(
            'unit_id,predicted_label\na,1\nb,1\nc,0\nd,0\n',
            'unit_id,reference_label\na,1\nb,0\nc,1\nd,0\n',
            'case_code,status\nT1,passed\nT2,failed\nT3,blocked\nT4,not_run\n',
            'synthetic_demo', 'Synthetic fixture', 'Four fixed grid cells',
        ), data_only=True)
        self.assertEqual(workbook.sheetnames, ['Summary', 'Classifications', 'Functional tests'])
        summary = workbook['Summary']
        self.assertEqual(summary['A2'].value, 'PPALMA Hotspot Evaluation')
        self.assertIn('DEMONSTRATION ONLY', summary['A3'].value)
        self.assertEqual(summary['B13'].value, .5)
        self.assertEqual(summary['B14'].value, .25)
        self.assertEqual(summary['B15'].value, .475)
        self.assertEqual(summary['B20'].value, 1)
        self.assertEqual(summary['C20'].value, 1)
        self.assertEqual(summary['B21'].value, 1)
        self.assertEqual(summary['C21'].value, 1)
        self.assertEqual(workbook['Classifications'].max_row, 9)
        self.assertEqual(workbook['Functional tests'].max_row, 9)


if __name__ == '__main__':
    unittest.main()
